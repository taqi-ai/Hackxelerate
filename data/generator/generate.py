"""Ascora synthetic campus data generator.

Builds a cohort of ~1,200 B.Tech students from hidden latent factors, then writes
seven deliberately messy source files (as a real college's ERP, attendance register,
LMS, clubs, placement cell, skills assessment and feedback system would produce).

Hidden truth (latent factors, outcomes, true identity of every source row) is written
to data/raw/_truth*.parquet. The app never reads it; it is used only to evaluate
entity resolution and the risk models honestly.

Usage:  python generate.py [--seed 42] [--out ../raw] [--n 1200]
"""
from __future__ import annotations

import argparse
import json
import random
import string
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from faker import Faker
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

# --------------------------------------------------------------------------- config

BRANCHES = {  # code: (display name, share of cohort, cs_family)
    "CSE": ("CSE", 0.26, True),
    "AIML": ("CSE-AIML", 0.20, True),
    "DS": ("CSE-Data Science", 0.16, True),
    "ECE": ("ECE", 0.20, False),
    "ME": ("ME", 0.18, False),
}
YEARS = {2: 25, 3: 24}  # current year of study -> batch (admission year, 2 digits)
CURRENT_SEM = {2: 3, 3: 5}
WEEKS_SO_FAR = 12  # semester week we are in (mid-sem exam was week 8)
SEM_START = date(2026, 7, 20)
EMAIL_DOMAIN = "cumail.in"

CS_CURRICULUM = {
    1: ["Mathematics-I", "Engineering Physics", "Programming in C", "Basic Electrical Engineering", "Engineering Graphics"],
    2: ["Mathematics-II", "Engineering Chemistry", "Object Oriented Programming", "Digital Electronics", "Communication Skills"],
    3: ["Discrete Mathematics", "Data Structures", "Computer Organization", "Database Management Systems", "Java Programming"],
    4: ["Design & Analysis of Algorithms", "Operating Systems", "Probability & Statistics", "Software Engineering", "Web Technologies"],
    5: ["Computer Networks", "Theory of Computation", "Machine Learning", "Compiler Design", "Cloud Computing"],
}
ECE_CURRICULUM = {
    1: CS_CURRICULUM[1],
    2: ["Mathematics-II", "Engineering Chemistry", "Electronic Devices", "Network Analysis", "Communication Skills"],
    3: ["Signals & Systems", "Analog Circuits", "Digital System Design", "Electromagnetic Theory", "Mathematics-III"],
    4: ["Control Systems", "Analog Communication", "Microprocessors", "Linear ICs", "Probability & Random Processes"],
    5: ["Digital Communication", "VLSI Design", "Digital Signal Processing", "Antennas & Wave Propagation", "Embedded Systems"],
}
ME_CURRICULUM = {
    1: CS_CURRICULUM[1],
    2: ["Mathematics-II", "Engineering Chemistry", "Engineering Mechanics", "Workshop Practice", "Communication Skills"],
    3: ["Thermodynamics", "Strength of Materials", "Manufacturing Processes", "Fluid Mechanics", "Mathematics-III"],
    4: ["Kinematics of Machines", "Applied Thermodynamics", "Machine Drawing", "Material Science", "Numerical Methods"],
    5: ["Heat Transfer", "Dynamics of Machines", "Design of Machine Elements", "IC Engines", "Metrology"],
}
SPECIALISATION_SEM5 = {
    "AIML": ["Computer Networks", "Deep Learning", "Machine Learning", "Natural Language Processing", "Cloud Computing"],
    "DS": ["Computer Networks", "Big Data Analytics", "Machine Learning", "Data Visualization", "Cloud Computing"],
}
MATH_SUBJECTS = {"Mathematics-I", "Mathematics-II", "Mathematics-III", "Discrete Mathematics",
                 "Probability & Statistics", "Probability & Random Processes", "Numerical Methods"}
# planted "gatekeeper" courses: weak performance predicts later backlogs
GATEKEEPERS = {"Mathematics-II", "Data Structures", "Signals & Systems", "Thermodynamics"}
CORE_LABS = {"Data Structures", "Database Management Systems", "Operating Systems", "Computer Networks",
             "Analog Circuits", "Digital System Design", "VLSI Design", "Thermodynamics", "Fluid Mechanics", "Heat Transfer"}

INTENTS = ["placement", "higher_studies", "govt_exams", "entrepreneurship", "undecided"]
INTENT_P = [0.62, 0.13, 0.10, 0.05, 0.10]

EVENTS = {
    "hackathon": ["Smart India Hackathon (internal round)", "HackCU 3.0", "byteXL CodeSprint", "Google Solution Challenge prep", "IEEE Hack Night"],
    "club": ["Coding Club", "Robotics Society", "Literary Society", "Entrepreneurship Cell", "NSS", "Music Club", "Dance Club", "Photography Club"],
    "workshop": ["Git & GitHub Workshop", "Resume Building Workshop", "IoT Bootcamp", "Cloud Fundamentals", "Public Speaking Masterclass", "CAD Workshop"],
    "certification": ["NPTEL: Python for Data Science", "AWS Cloud Practitioner", "Coursera: Machine Learning", "NPTEL: DBMS", "Cisco CCNA Module 1", "Autodesk Fusion 360"],
}

FEEDBACK_TEMPLATES = {
    "en": {
        "pos": ["The professor explains concepts clearly and the labs are useful.", "Good course, assignments helped me understand the topic.",
                "Very engaging classes and doubts are cleared patiently.", "Well structured course with practical examples."],
        "neu": ["Course is okay but the pace is a bit fast.", "Content is fine, would like more practice problems.",
                "Average experience, lab equipment could be better."],
        "neg": ["Too much syllabus and very little time, hard to keep up.", "Classes are boring and mostly reading from slides.",
                "Assignments are too many and deadlines clash with other subjects.", "Lab sessions are often cancelled, I feel lost."],
    },
    "hi": {
        "pos": ["सर बहुत अच्छे से पढ़ाते हैं, लैब भी उपयोगी है।", "कोर्स अच्छा है, असाइनमेंट से समझ बढ़ी।"],
        "neu": ["कोर्स ठीक है लेकिन पढ़ाने की गति थोड़ी तेज़ है।", "सामग्री ठीक है, और अभ्यास प्रश्न चाहिए।"],
        "neg": ["सिलेबस बहुत ज़्यादा है और समय बहुत कम।", "क्लास में समझ नहीं आता, सिर्फ स्लाइड पढ़ी जाती हैं।", "लैब अक्सर रद्द हो जाती है।"],
    },
    "hinglish": {
        "pos": ["Sir ka teaching style bahut accha hai, concepts clear ho jaate hain.", "Labs mast hain, practical knowledge mil raha hai."],
        "neu": ["Course theek hai but pace thoda fast hai.", "Content ok hai, thode aur examples chahiye."],
        "neg": ["Bahut zyada assignments hain, deadlines clash ho rahi hain.", "Class mein kuch samajh nahi aata, sirf slides padhte hain.",
                "Lab cancel hoti rehti hai, practical bilkul nahi ho raha."],
    },
}


def sigmoid(x):
    return 1.0 / (1.0 + np.exp(-x))


def clip(x, lo, hi):
    return float(min(hi, max(lo, x)))


@dataclass
class Ctx:
    rng: np.random.Generator
    py: random.Random
    fake: Faker
    out: Path
    report: dict


# --------------------------------------------------------------------------- cohort

def curriculum(branch: str, sem: int) -> list[str]:
    if sem == 5 and branch in SPECIALISATION_SEM5:
        return SPECIALISATION_SEM5[branch]
    if branch in ("CSE", "AIML", "DS"):
        return CS_CURRICULUM[sem]
    return ECE_CURRICULUM[sem] if branch == "ECE" else ME_CURRICULUM[sem]


def make_cohort(ctx: Ctx, n: int) -> pd.DataFrame:
    rng, fake = ctx.rng, ctx.fake
    rows = []
    codes = list(BRANCHES)
    shares = np.array([BRANCHES[c][1] for c in codes])
    seq = {}
    for i in range(n):
        year = 2 if rng.random() < 0.52 else 3
        branch = codes[rng.choice(len(codes), p=shares / shares.sum())]
        batch = YEARS[year]
        key = (batch, branch)
        seq[key] = seq.get(key, 0) + 1
        gender = "F" if rng.random() < 0.38 else "M"
        first = fake.first_name_female() if gender == "F" else fake.first_name_male()
        last = fake.last_name()
        hosteller = bool(rng.random() < 0.45)
        first_gen = bool(rng.random() < 0.25)
        lateral = bool(year == 3 and rng.random() < 0.08)
        medium = "Hindi" if rng.random() < 0.30 else "English"
        intent = INTENTS[rng.choice(len(INTENTS), p=INTENT_P)]
        rows.append(dict(
            student_id=f"CU{batch}{branch}{seq[key]:04d}", seq=seq[key], batch=batch, branch_code=branch,
            branch=BRANCHES[branch][0], cs_family=BRANCHES[branch][2], year=year, current_sem=CURRENT_SEM[year],
            first_name=first, last_name=last, name=f"{first} {last}", gender=gender,
            hosteller=hosteller, first_generation=first_gen, lateral_entry=lateral, schooling_medium=medium,
            career_intent=intent,
            # latent factors
            ability=rng.normal(), math_aptitude=rng.normal(), motivation=rng.normal(), consistency=rng.normal(),
            financial_stress=rng.beta(2, 5), soft_base=rng.normal(), extracurricular=rng.normal(),
            satisfaction=rng.normal(),
            phone="9" + "".join(str(d) for d in rng.integers(0, 10, 9)),
        ))
    df = pd.DataFrame(rows)

    # mediated context effects (small, realistic): commute -> attendance, exposure -> soft skills
    df["commute_penalty"] = np.where(df.hosteller, 0.0, rng.uniform(0, 6, len(df)))
    df["soft"] = df.soft_base + 0.35 * df.motivation - 0.25 * (df.schooling_medium == "Hindi") - 0.12 * df.first_generation
    df["prep"] = 0.7 * df.motivation + 0.6 * (df.career_intent == "placement") + rng.normal(0, 0.6, len(df))

    # planted subgroups
    df["silent_decliner"] = rng.random(len(df)) < 0.08
    # decliners look average today: pull extreme latents toward the middle
    df.loc[df.silent_decliner, "ability"] = rng.normal(0.15, 0.45, df.silent_decliner.sum())
    df["engaged_slipping"] = (~df.silent_decliner) & (rng.random(len(df)) < 0.06)
    df.loc[df.engaged_slipping, "extracurricular"] = rng.normal(1.6, 0.4, df.engaged_slipping.sum())
    df.loc[df.engaged_slipping, "consistency"] = rng.normal(-1.0, 0.4, df.engaged_slipping.sum())
    df["post_midsem_dropoff"] = (~df.silent_decliner) & (rng.random(len(df)) < sigmoid(-2.6 - 0.8 * df.motivation - 0.5 * df.consistency + 1.2 * df.financial_stress))

    # mentors: ~50 students each, by branch + year + section
    df["section"] = ""
    mentor_rows = []
    m = 0
    for (branch, year), g in df.groupby(["branch_code", "year"], sort=True):
        n_sections = max(1, round(len(g) / 60))
        idx = g.sort_values("seq").index
        for s, chunk in enumerate(np.array_split(idx, n_sections)):
            df.loc[chunk, "section"] = "ABCD"[s]
            n_m = max(1, round(len(chunk) / 50))
            for part in np.array_split(chunk, n_m):
                m += 1
                title = "Dr." if rng.random() < 0.6 else "Prof."
                mname = f"{title} {fake.first_name()} {fake.last_name()}"
                dept = "CSE" if BRANCHES[branch][2] else BRANCHES[branch][0]
                mentor_rows.append(dict(id=f"M{m:02d}", name=mname, department=dept))
                df.loc[part, "mentor_id"] = f"M{m:02d}"
    ctx.mentors = pd.DataFrame(mentor_rows)
    return df


# --------------------------------------------------------------------------- observables

def grade_point(total: float) -> int:
    for cut, gp in ((90, 10), (80, 9), (70, 8), (60, 7), (50, 6), (45, 5), (40, 4)):
        if total >= cut:
            return gp
    return 0


def make_academics(ctx: Ctx, df: pd.DataFrame) -> pd.DataFrame:
    rng = ctx.rng
    recs = []
    for s in df.itertuples():
        for sem in range(1, s.current_sem + 1):
            if s.lateral_entry and sem <= 2:
                continue  # lateral entrants join in year 2
            current = sem == s.current_sem
            for subj in curriculum(s.branch_code, sem):
                mu = 61 + 10 * s.ability + 4 * s.motivation + 3 * s.consistency
                if subj in MATH_SUBJECTS:
                    mu += 8 * s.math_aptitude
                if subj in GATEKEEPERS:
                    mu -= 4  # harder course
                if s.engaged_slipping and sem >= s.current_sem - 1:
                    mu -= 7
                total_latent = mu + rng.normal(0, 8)
                internal = clip(total_latent * 0.4 + rng.normal(0, 3), 4, 40)
                if current:
                    # current semester: internals so far (mid-sem week 8), no end-sem yet
                    if s.silent_decliner:
                        internal = clip(internal - rng.uniform(4, 8), 3, 40)
                    if s.post_midsem_dropoff:
                        internal = clip(internal - rng.uniform(2, 5), 3, 40)
                    recs.append(dict(student_id=s.student_id, sem=sem, subject=subj, internal=round(internal, 1),
                                     end_sem=np.nan, total=np.nan, current=True))
                else:
                    end_sem = clip(total_latent * 0.6 + rng.normal(0, 5), 0, 60)
                    total = internal + end_sem
                    recs.append(dict(student_id=s.student_id, sem=sem, subject=subj, internal=round(internal, 1),
                                     end_sem=round(end_sem, 1), total=round(total, 1), current=False))
    marks = pd.DataFrame(recs)
    past = marks[~marks.current].copy()
    past["gp"] = past.total.apply(grade_point)
    past["failed"] = (past.end_sem < 24) | (past.total < 40)
    sgpa = past.groupby(["student_id", "sem"]).gp.mean().round(2).rename("sgpa").reset_index()
    ctx.sgpa = sgpa
    # active backlogs: failed and not yet cleared (60% get cleared in supplementary exams)
    past["cleared"] = ctx.rng.random(len(past)) < 0.6
    backlogs = past[past.failed & ~past.cleared].groupby("student_id").size()
    df["backlogs_true"] = df.student_id.map(backlogs).fillna(0).astype(int)
    df["cgpa_true"] = df.student_id.map(sgpa.groupby("student_id").sgpa.mean().round(2))
    gk = past[past.subject.isin(GATEKEEPERS)].groupby("student_id").total.mean()
    df["gatekeeper_total"] = df.student_id.map(gk)
    return marks


def make_weekly(ctx: Ctx, df: pd.DataFrame) -> pd.DataFrame:
    rng = ctx.rng
    rows = []
    for s in df.itertuples():
        base_login = max(0.5, 4.2 + 1.4 * s.motivation + 0.6 * s.consistency + rng.normal(0, 0.6))
        base_comp = sigmoid(1.4 + 0.9 * s.motivation + 0.7 * s.consistency - 0.4 * s.extracurricular * s.engaged_slipping)
        base_att = 82 + 5 * s.motivation + 4 * s.consistency - s.commute_penalty - 6 * s.financial_stress
        for w in range(1, WEEKS_SO_FAR + 1):
            decay = 1.0
            if s.silent_decliner and w > 5:
                decay = max(0.35, 1 - 0.075 * (w - 5))
            if s.post_midsem_dropoff and w > 8:
                decay = min(decay, max(0.3, 1 - 0.17 * (w - 8)))
            logins = max(0, rng.poisson(base_login * decay))
            due = int(rng.integers(1, 4))
            submitted = int(min(due, rng.binomial(due, clip(base_comp * (0.55 + 0.45 * decay), 0.02, 0.99))))
            minutes = max(0, rng.normal(38 * base_login * decay, 25))
            att = clip(base_att * (0.75 + 0.25 * decay) + rng.normal(0, 6), 10, 100)
            rows.append(dict(student_id=s.student_id, week=w, logins=logins, due=due, submitted=submitted,
                             minutes=round(minutes), attendance=att))
    return pd.DataFrame(rows)


def make_outcomes(ctx: Ctx, df: pd.DataFrame, weekly: pd.DataFrame) -> None:
    rng = ctx.rng
    late = weekly[weekly.week > 8].groupby("student_id").attendance.mean()
    att = df.student_id.map(late)
    gk = (df.gatekeeper_total.fillna(df.gatekeeper_total.median()) - 60) / 12
    logit = (-2.95 - 1.15 * df.ability - 0.55 * df.motivation - 0.45 * df.consistency + 0.9 * df.financial_stress
             + 1.5 * df.silent_decliner + 1.2 * df.post_midsem_dropoff + 0.9 * df.engaged_slipping
             - 0.45 * gk - 0.035 * (att - 75) + 0.35 * np.minimum(df.backlogs_true, 3) + rng.normal(0, 0.5, len(df)))
    df["p_academic_fail"] = sigmoid(logit)
    df["end_sem_fail_or_backlog"] = rng.random(len(df)) < df.p_academic_fail
    pl = (0.75 + 0.45 * df.ability + 0.95 * df.soft + 0.85 * df.prep - 1.0 * df.silent_decliner
          + 0.35 * df.cs_family + rng.normal(0, 0.5, len(df)))
    df["p_placed"] = sigmoid(pl)
    placed = rng.random(len(df)) < df.p_placed
    df["placed_by_final_year"] = np.where(df.career_intent == "placement", placed, None)


# --------------------------------------------------------------------------- messy writers

def roll_erp(s) -> str:
    return s.student_id  # CU25AIML0042


def roll_register(s) -> str:
    return f"{s.batch}-{s.branch_code}-{s.seq}"  # 25-AIML-42


def roll_placement(s) -> str:
    return f"20{s.batch}/{s.branch_code}/{s.seq:03d}"  # 2025/AIML/042


def typo(word: str, py: random.Random) -> str:
    if len(word) < 4:
        return word
    i = py.randrange(1, len(word) - 1)
    op = py.choice(["swap", "drop", "double"])
    if op == "swap":
        return word[:i] + word[i + 1] + word[i] + word[i + 2:]
    if op == "drop":
        return word[:i] + word[i + 1:]
    return word[:i] + word[i] + word[i:]


def fmt_date(d: date, py: random.Random) -> str:
    return py.choice([d.isoformat(), d.strftime("%d/%m/%Y"), d.strftime("%b %d, %Y")])


def write_erp(ctx: Ctx, df: pd.DataFrame, marks: pd.DataFrame, weekly: pd.DataFrame, truth_links: list) -> None:
    rng, py = ctx.rng, ctx.py
    sgpa = ctx.sgpa.pivot(index="student_id", columns="sem", values="sgpa")
    att_erp = weekly.groupby("student_id").attendance.mean()
    by_id = df.set_index("student_id")
    rows = []
    conflicts_att = set(py.sample(list(df.student_id), int(0.03 * len(df))))
    branch_conflicts = set(py.sample(list(df.student_id), int(0.01 * len(df))))
    for r in marks.itertuples():
        s = by_id.loc[r.student_id]
        a = att_erp[r.student_id]
        if r.student_id in conflicts_att:
            a = a + py.choice([-1, 1]) * rng.uniform(7, 15)  # ERP attendance disagrees with register
        rows.append({
            "Roll No": r.student_id, "Student Name": s["name"],
            "Branch": s.branch if r.student_id not in branch_conflicts else py.choice([b[0] for b in BRANCHES.values() if b[0] != s.branch]),
            "Semester": r.sem, "Subject": r.subject, "Internal (40)": r.internal,
            "End Sem (60)": r.end_sem, "Total (100)": r.total,
            "SGPA": sgpa.loc[r.student_id].get(r.sem, np.nan) if r.sem in sgpa.columns else np.nan,
            "CGPA": s.cgpa_true, "Active Backlogs": s.backlogs_true,
            "Attendance %": round(clip(a, 0, 100), 1) if r.current else np.nan,
        })
    erp = pd.DataFrame(rows)
    n = len(erp)
    # planted messiness
    miss = rng.random(n) < 0.05
    erp.loc[miss, "Internal (40)"] = np.nan
    oor = erp.sample(12, random_state=1).index
    erp.loc[oor, "Internal (40)"] = rng.choice([45, 52, -3, 99], len(oor))  # out of range
    dup_students = py.sample(list(df.student_id), int(0.03 * len(df)))
    dups = erp[erp["Roll No"].isin(dup_students) & (erp["Semester"] == erp.groupby("Roll No")["Semester"].transform("max"))].copy()
    dups["Student Name"] = dups["Student Name"].str.upper() + " "  # same student, re-entered
    erp = pd.concat([erp, dups]).sample(frac=1, random_state=7).reset_index(drop=True)
    erp.to_csv(ctx.out / "erp_academics.csv", index=False)
    for i, rid in enumerate(erp["Roll No"]):
        truth_links.append(("erp_academics", f"row {i + 2}", rid))
    ctx.report["erp_academics"] = dict(rows=len(erp), duplicate_rows=len(dups), missing_internal=int(miss.sum()),
                                       out_of_range=len(oor), attendance_conflicts=len(conflicts_att),
                                       branch_conflicts=len(branch_conflicts))


def write_attendance(ctx: Ctx, df: pd.DataFrame, weekly: pd.DataFrame, truth_links: list) -> None:
    rng = ctx.rng
    months = {"Aug": range(1, 5), "Sep": range(5, 9), "Oct": range(9, 13)}
    wb = Workbook()
    wb.remove(wb.active)
    missing_total = 0
    by_id = df.set_index("student_id")
    for mon, weeks in months.items():
        ws = wb.create_sheet(f"{mon} 2026")
        ws.merge_cells("A1:E1")
        ws["A1"] = f"ATTENDANCE REGISTER - {mon.upper()} 2026 - B.TECH (ALL BRANCHES)"
        ws["A1"].font = Font(bold=True, size=13)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws.merge_cells("A2:E2")
        ws["A2"] = "Generated by Academic Section. Classes held / attended per subject."
        ws.append(["Roll Number", "Subject", "Classes Held", "Classes Attended", "Remarks"])
        for c in ws[3]:
            c.font = Font(bold=True)
        wk = weekly[weekly.week.isin(list(weeks))].groupby("student_id").attendance.mean()
        skip = set(rng.choice(df.student_id, int(0.04 * len(df)), replace=False))  # missing from this month's sheet
        missing_total += len(skip)
        row_no = 4
        for sid in df.student_id:
            if sid in skip:
                continue
            s = by_id.loc[sid]
            for subj in curriculum(s.branch_code, s.current_sem):
                held = int(rng.integers(10, 15))
                pct = clip(wk[sid] + rng.normal(0, 7) - (4 if subj in CORE_LABS and s.silent_decliner else 0), 0, 100)
                attended = int(round(held * pct / 100))
                if rng.random() < 0.002:
                    attended = held + int(rng.integers(1, 5))  # impossible value
                roll = roll_register(s)
                ws.append([roll, subj, held, attended, ""])
                truth_links.append(("attendance_register", f"{mon} 2026 row {row_no}", sid))
                row_no += 1
    wb.save(ctx.out / "attendance_register.xlsx")
    ctx.report["attendance_register"] = dict(sheets=list(months), student_month_gaps=missing_total,
                                             merged_header_rows=2)


def email_for(s, py: random.Random, variant: bool) -> str:
    f, l = s.first_name.lower(), s.last_name.lower()
    base = f"{f}.{l}"
    if variant:
        base = py.choice([f"{f[0]}.{l}", f"{typo(f, py)}.{l}", f"{f}.{l}{py.randint(1, 9)}", f"{f}{l}"])
    return f"{base}@{EMAIL_DOMAIN}"


def write_lms(ctx: Ctx, df: pd.DataFrame, weekly: pd.DataFrame, truth_links: list) -> None:
    py = ctx.py
    by_id = df.set_index("student_id")
    primary = {}
    variant_ids = set(py.sample(list(df.student_id), int(0.12 * len(df))))
    second_account = set(py.sample(list(df.student_id), int(0.02 * len(df))))  # same student, two LMS logins
    for sid in df.student_id:
        primary[sid] = email_for(by_id.loc[sid], py, sid in variant_ids)
    rows = []
    for r in weekly.itertuples():
        s = by_id.loc[r.student_id]
        email = primary[r.student_id]
        if r.student_id in second_account and r.week % 3 == 0:
            email = email_for(s, py, True)
        wk_date = SEM_START + timedelta(weeks=r.week - 1)
        rows.append({"user_email": email, "program": f"B.Tech {s.branch} {2000 + s.batch}",
                     "week_start": wk_date.isoformat() if py.random() < 0.6 else wk_date.strftime("%d-%m-%Y"),
                     "logins": r.logins, "assignments_due": r.due, "assignments_submitted": r.submitted,
                     "minutes_active": r.minutes})
        truth_links.append(("lms_activity_log", email, r.student_id))
    lms = pd.DataFrame(rows)
    miss = ctx.rng.random(len(lms)) < 0.04
    lms.loc[miss, "minutes_active"] = np.nan
    lms.to_csv(ctx.out / "lms_activity_log.csv", index=False)
    ctx.report["lms_activity_log"] = dict(rows=len(lms), email_variants=len(variant_ids),
                                          students_with_two_accounts=len(second_account), date_formats=2,
                                          missing_minutes=int(miss.sum()))


def name_variant(s, py: random.Random) -> str:
    r = py.random()
    if r < 0.55:
        return s.name
    if r < 0.70:
        return f"{s.first_name[0]}. {s.last_name}"
    if r < 0.80:
        return f"{s.last_name} {s.first_name}"
    if r < 0.90:
        return f"{typo(s.first_name, py)} {s.last_name}"
    return s.name.upper() if py.random() < 0.5 else s.name.lower()


def write_clubs(ctx: Ctx, df: pd.DataFrame, truth_links: list) -> pd.DataFrame:
    rng, py = ctx.rng, ctx.py
    rows = []
    for s in df.itertuples():
        lam = {
            "hackathon": max(0.02, 0.35 + 0.5 * s.motivation + 0.45 * s.cs_family + 0.6 * s.extracurricular * (s.extracurricular > 0)),
            "club": max(0.05, 0.9 + 0.4 * s.extracurricular + 0.3 * s.hosteller),
            "workshop": max(0.05, 0.9 + 0.4 * s.motivation),
            "certification": max(0.02, 0.6 + 0.5 * s.motivation + 0.2 * s.prep),
        }
        for etype, l in lam.items():
            for _ in range(rng.poisson(l)):
                d = date(2025, 8, 1) + timedelta(days=int(rng.integers(0, 440)))
                rows.append({"Participant Name": name_variant(s, py), "Branch": s.branch, "Year": s.year,
                             "Event / Activity": py.choice(EVENTS[etype]), "Type": etype.capitalize(),
                             "Date": fmt_date(d, py), "_sid": s.student_id})
    clubs = pd.DataFrame(rows).sample(frac=1, random_state=3).reset_index(drop=True)
    for i, sid in enumerate(clubs._sid):
        truth_links.append(("clubs_events", f"row {i + 2}", sid))
    clubs.drop(columns="_sid").to_csv(ctx.out / "clubs_events.csv", index=False)
    ctx.report["clubs_events"] = dict(rows=len(clubs), identifier="name only", date_formats=3,
                                      initials_or_variants=int((clubs["Participant Name"].str.contains(r"\.")).sum()))
    return clubs


def write_placement(ctx: Ctx, df: pd.DataFrame, truth_links: list) -> None:
    rng, py = ctx.rng, ctx.py
    rows = []
    for s in df.itertuples():
        visited = rng.random() < (0.62 if s.year == 2 else 0.9)
        if not visited:
            rows.append({"Roll": roll_placement(s), "Name": s.name, "Branch": s.branch, "Contact": s.phone,
                         "Aptitude Score": None, "Coding Score": None, "Mock Interviews Attended": None,
                         "Avg Mock Rating (1-5)": None, "Career Intent (declared)": None})
            continue
        apt = clip(55 + 10 * s.ability + 6 * s.motivation + rng.normal(0, 9), 0, 100)
        cod = clip(48 + 11 * s.ability + 9 * s.prep + 9 * s.cs_family + rng.normal(0, 10), 0, 100)
        mocks = int(rng.poisson(max(0.05, (1.3 + 1.3 * s.prep) * (0.55 if s.year == 2 else 1.0))))
        rating = round(clip(2.8 + 0.65 * s.soft + 0.25 * s.ability + rng.normal(0, 0.4), 1, 5), 1) if mocks else None
        declared = s.career_intent if rng.random() < 0.9 else "undecided"
        rows.append({"Roll": roll_placement(s), "Name": s.name, "Branch": s.branch, "Contact": s.phone,
                     "Aptitude Score": round(apt, 1), "Coding Score": round(cod, 1),
                     "Mock Interviews Attended": mocks, "Avg Mock Rating (1-5)": rating,
                     "Career Intent (declared)": declared.replace("_", " ").title()})
    pl = pd.DataFrame(rows)
    dup = pl.sample(int(0.03 * len(pl)), random_state=5)
    pl = pd.concat([pl, dup]).sample(frac=1, random_state=11).reset_index(drop=True)
    oor = pl.dropna(subset=["Aptitude Score"]).sample(6, random_state=2).index
    pl.loc[oor, "Aptitude Score"] = [112, -5, 140, 101, -1, 250]
    by_roll = {roll_placement(s): s.student_id for s in df.itertuples()}
    for i, r in enumerate(pl.Roll):
        truth_links.append(("placement_cell", f"row {i + 2}", by_roll[r]))
    with pd.ExcelWriter(ctx.out / "placement_cell.xlsx", engine="openpyxl") as xw:
        pl.to_excel(xw, sheet_name="Placement Readiness", index=False, startrow=1)
        xw.sheets["Placement Readiness"]["A1"] = "Training & Placement Cell - Readiness Tracker (2026-27)"
    ctx.report["placement_cell"] = dict(rows=len(pl), duplicate_rows=len(dup),
                                        never_visited=int(pl["Aptitude Score"].isna().sum()), out_of_range=len(oor))


def write_skills(ctx: Ctx, df: pd.DataFrame, truth_links: list) -> None:
    rng = ctx.rng
    records = []
    ten_scale = 0
    for s in df.itertuples():
        if rng.random() < 0.05:
            continue  # absent for the assessment
        tech = {k: clip(55 + 11 * s.ability + 6 * s.prep + off + rng.normal(0, 9), 0, 100)
                for k, off in (("programming", 6 * s.cs_family), ("problem_solving_dsa", 0), ("domain_knowledge", 2))}
        soft = {k: clip(58 + 12 * s.soft + rng.normal(0, 8) - (5 if k == "communication" and s.schooling_medium == "Hindi" else 0), 0, 100)
                for k in ("communication", "teamwork", "problem_solving", "presentation")}
        scale10 = rng.random() < 0.15
        if scale10:
            ten_scale += 1
            tech = {k: round(v / 10, 1) for k, v in tech.items()}
            soft = {k: round(v / 10, 1) for k, v in soft.items()}
        else:
            tech = {k: round(v) for k, v in tech.items()}
            soft = {k: round(v) for k, v in soft.items()}
        rid = s.student_id.lower() if rng.random() < 0.5 else s.student_id
        records.append({"candidate": {"roll": rid, "name": s.name},
                        "assessment": {"cycle": "2026-ODD", "vendor": "SkillLens" if scale10 else "CampusAssess",
                                       "technical": tech, "soft_skills": soft}})
        truth_links.append(("skills_assessment", rid, s.student_id))
    payload = {"exported_at": "2026-10-09T18:30:00+05:30", "count": len(records), "records": records}
    (ctx.out / "skills_assessment.json").write_text(json.dumps(payload, ensure_ascii=False, indent=1), encoding="utf-8")
    ctx.report["skills_assessment"] = dict(records=len(records), on_0_10_scale=ten_scale, nested=True)


def write_feedback(ctx: Ctx, df: pd.DataFrame, truth_links: list) -> None:
    rng, py = ctx.rng, ctx.py
    rows = []
    for s in df.itertuples():
        for _ in range(rng.binomial(2, 0.55)):
            sat = s.satisfaction + 0.3 * s.motivation - 0.6 * s.silent_decliner + rng.normal(0, 0.7)
            rating = int(np.clip(round(3.2 + 0.9 * sat), 1, 5))
            tone = "pos" if rating >= 4 else ("neu" if rating == 3 else "neg")
            lang = py.choices(["en", "hi", "hinglish"], weights=[0.55, 0.15, 0.30])[0]
            if s.schooling_medium == "Hindi" and py.random() < 0.3:
                lang = "hi"
            anon = rng.random() < 0.3
            subj = py.choice(curriculum(s.branch_code, s.current_sem))
            rows.append({"roll_no": "" if anon else s.student_id, "course": subj, "rating": rating,
                         "comment": py.choice(FEEDBACK_TEMPLATES[lang][tone]), "submitted_on": fmt_date(date(2026, 9, 25) + timedelta(days=int(rng.integers(0, 12))), py),
                         "_sid": s.student_id, "_lang": lang})
    fb = pd.DataFrame(rows).sample(frac=1, random_state=9).reset_index(drop=True)
    for i, sid in enumerate(fb._sid):
        truth_links.append(("feedback_forms", f"row {i + 2}", sid))
    fb.drop(columns=["_sid", "_lang"]).to_csv(ctx.out / "feedback_forms.csv", index=False, encoding="utf-8")
    ctx.report["feedback_forms"] = dict(rows=len(fb), anonymous=int((fb.roll_no == "").sum()),
                                        languages=fb._lang.value_counts().to_dict())


def write_intervention_history(ctx: Ctx, df: pd.DataFrame) -> None:
    """Last semester's intervention records (used by the 'What works' analysis).

    Effects are planted per type so the comparison has a known right answer."""
    rng = ctx.rng
    effect = {"mentor_meeting": 0.06, "peer_tutoring": 0.10, "mock_interview": 0.12, "counselling_referral": 0.04,
              "parent_call": 0.02, "skill_workshop": 0.07}
    rows = []
    cand = df[df.p_academic_fail > 0.3].sample(frac=0.6, random_state=4)
    for s in cand.itertuples():
        if rng.random() < 0.55:
            t = ctx.py.choice(list(effect))
            improved = rng.random() < (0.30 + effect[t] * 2.2)
            rows.append(dict(student_id=s.student_id, semester=s.current_sem - 1, type=t, received=True, improved=improved))
        else:
            rows.append(dict(student_id=s.student_id, semester=s.current_sem - 1, type=None, received=False,
                             improved=rng.random() < 0.30))
    pd.DataFrame(rows).to_csv(ctx.out / "_intervention_history.csv", index=False)


# --------------------------------------------------------------------------- main

def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=42)
    ap.add_argument("--n", type=int, default=1200)
    ap.add_argument("--out", type=Path, default=Path(__file__).resolve().parents[1] / "raw")
    args = ap.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)

    rng = np.random.default_rng(args.seed)
    py = random.Random(args.seed)
    fake = Faker("en_IN")
    Faker.seed(args.seed)
    ctx = Ctx(rng=rng, py=py, fake=fake, out=args.out, report={})

    df = make_cohort(ctx, args.n)
    marks = make_academics(ctx, df)
    weekly = make_weekly(ctx, df)
    make_outcomes(ctx, df, weekly)

    links: list = []
    write_erp(ctx, df, marks, weekly, links)
    write_attendance(ctx, df, weekly, links)
    write_lms(ctx, df, weekly, links)
    write_clubs(ctx, df, links)
    write_placement(ctx, df, links)
    write_skills(ctx, df, links)
    write_feedback(ctx, df, links)
    write_intervention_history(ctx, df)

    ctx.mentors.to_csv(args.out / "mentors.csv", index=False)
    truth_cols = ["student_id", "name", "branch", "year", "section", "mentor_id", "gender", "hosteller", "first_generation",
                  "lateral_entry", "schooling_medium", "career_intent", "ability", "math_aptitude", "motivation", "consistency",
                  "financial_stress", "soft", "prep", "extracurricular", "silent_decliner", "engaged_slipping",
                  "post_midsem_dropoff", "p_academic_fail", "end_sem_fail_or_backlog", "p_placed", "placed_by_final_year"]
    truth = df[truth_cols].copy()
    truth["placed_by_final_year"] = truth.placed_by_final_year.astype("boolean")
    truth.to_parquet(args.out / "_truth.parquet", index=False)
    pd.DataFrame(links, columns=["source", "raw_ref", "student_id"]).drop_duplicates().to_parquet(
        args.out / "_truth_links.parquet", index=False)

    ctx.report["cohort"] = dict(
        seed=args.seed, students=len(df), mentors=len(ctx.mentors),
        by_branch=df.branch.value_counts().to_dict(), by_year=df.year.value_counts().to_dict(),
        silent_decliners=int(df.silent_decliner.sum()), engaged_but_slipping=int(df.engaged_slipping.sum()),
        post_midsem_dropoff=int(df.post_midsem_dropoff.sum()),
        academic_fail_rate=round(float(df.end_sem_fail_or_backlog.mean()), 3),
        placement_intent=int((df.career_intent == "placement").sum()),
        unplaced_rate_among_placement_intent=round(float(1 - df.loc[df.career_intent == "placement", "placed_by_final_year"].astype(bool).mean()), 3),
        planted_patterns=[
            "LMS activity decays after week 5 for silent decliners (~8%) whose current scores look average",
            "LMS + attendance drop-off after mid-sem (week 8) for a motivation/consistency/financial-stress driven group",
            "Gatekeeper courses (Mathematics-II, Data Structures, Signals & Systems, Thermodynamics) predict later failure",
            "High CGPA alone does not predict placement; soft skills and mock-interview preparation dominate",
            "Engaged-but-slipping group: high extracurricular activity, falling academics",
            "Context effects are mediated only: day scholars' commute lowers attendance; schooling medium affects communication scores",
        ],
    )
    (args.out / "generation_report.json").write_text(json.dumps(ctx.report, indent=2, default=str), encoding="utf-8")
    print(json.dumps(ctx.report["cohort"], indent=2, default=str))


if __name__ == "__main__":
    main()
