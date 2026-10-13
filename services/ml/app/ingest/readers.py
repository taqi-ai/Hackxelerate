"""File readers. Every output row keeps a `_ref` lineage pointer back to file + sheet + row."""
from __future__ import annotations

import json
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .schema import Mapping, detect_source, map_headers

SOURCE_FILES = {
    "erp_academics": "erp_academics.csv",
    "attendance_register": "attendance_register.xlsx",
    "lms_activity_log": "lms_activity_log.csv",
    "clubs_events": "clubs_events.csv",
    "placement_cell": "placement_cell.xlsx",
    "skills_assessment": "skills_assessment.json",
    "feedback_forms": "feedback_forms.csv",
}
SOURCE_LABELS = {
    "erp_academics": "Academic (ERP)",
    "attendance_register": "Attendance register",
    "lms_activity_log": "LMS activity",
    "clubs_events": "Clubs & events",
    "placement_cell": "Placement cell",
    "skills_assessment": "Skills assessment",
    "feedback_forms": "Feedback forms",
}


@dataclass
class SourceFrame:
    source: str
    file: str
    df: pd.DataFrame
    mapping: dict
    rows_in: int
    notes: list[str]


def _apply_mapping(df: pd.DataFrame, m: Mapping) -> pd.DataFrame:
    keep = {raw: canon for raw, canon in m.columns.items()}
    out = df[list(keep)].rename(columns=keep)
    return out


def _find_header_row(rows: list[tuple], max_scan: int = 10) -> int:
    """Title/merged rows sit above the real header: pick the first row that maps like a header."""
    best_i, best_score = 0, -1.0
    for i, row in enumerate(rows[:max_scan]):
        cells = [str(c) for c in row if c is not None and str(c).strip()]
        if len(cells) < 2:
            continue
        score = detect_source(cells).score
        if score > best_score:
            best_i, best_score = i, score
    return best_i


def read_csv(path: Path) -> SourceFrame:
    raw = pd.read_csv(path, dtype=str, keep_default_na=False, na_values=[""], encoding="utf-8")
    m = detect_source(list(raw.columns))
    df = _apply_mapping(raw, m)
    df["src_ref"] = [f"{path.name} row {i + 2}" for i in range(len(df))]
    df["src_key"] = [f"row {i + 2}" for i in range(len(df))]
    return SourceFrame(m.source_type, path.name, df, m.as_dict(), len(raw), [])


def read_xlsx(path: Path) -> SourceFrame:
    wb = load_workbook(path, read_only=True, data_only=True)
    frames, mapping, notes, total = [], None, [], 0
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        h = _find_header_row(rows)
        if h > 0:
            notes.append(f"Sheet '{ws.title}': skipped {h} title/merged row(s) above the header")
        headers = [str(c) if c is not None else f"col{j}" for j, c in enumerate(rows[h])]
        body = rows[h + 1:]
        raw = pd.DataFrame(body, columns=headers)
        raw = raw.dropna(how="all")
        m = detect_source(headers) if mapping is None else map_headers(headers, mapping.source_type)
        mapping = mapping or m
        df = _apply_mapping(raw, m)
        excel_rows = [h + 2 + int(i) for i in raw.index]  # 1-based excel row numbers
        df["src_ref"] = [f"{path.name}, sheet '{ws.title}', row {r}" for r in excel_rows]
        df["src_key"] = [f"{ws.title} row {r}" for r in excel_rows]
        df["src_sheet"] = ws.title
        frames.append(df)
        total += len(raw)
    wb.close()
    df = pd.concat(frames, ignore_index=True)
    return SourceFrame(mapping.source_type, path.name, df, mapping.as_dict(), total, notes)


def read_skills_json(path: Path) -> SourceFrame:
    payload = json.loads(path.read_text(encoding="utf-8"))
    recs = payload["records"] if isinstance(payload, dict) else payload
    rows = []
    for i, r in enumerate(recs):
        cand, asm = r.get("candidate", {}), r.get("assessment", {})
        rows.append({"roll": cand.get("roll"), "name": cand.get("name"), "vendor": asm.get("vendor"),
                     "technical": asm.get("technical", {}), "soft": asm.get("soft_skills", {}),
                     "src_ref": f"{path.name} records[{i}]", "src_key": cand.get("roll")})
    df = pd.DataFrame(rows)
    mapping = {"source_type": "skills_assessment",
               "columns": {"candidate.roll": "roll", "candidate.name": "name", "assessment.technical.*": "technical",
                           "assessment.soft_skills.*": "soft"},
               "method": {"roll": "json-path", "technical": "json-path", "soft": "json-path"}, "unmapped": [], "score": 1.0}
    return SourceFrame("skills_assessment", path.name, df, mapping, len(rows), ["Flattened nested JSON records"])


def read_any(path: Path) -> SourceFrame:
    suf = path.suffix.lower()
    if suf == ".json":
        return read_skills_json(path)
    if suf in (".xlsx", ".xls"):
        return read_xlsx(path)
    return read_csv(path)


def read_folder(folder: Path) -> dict[str, SourceFrame]:
    out: dict[str, SourceFrame] = {}
    for p in sorted(folder.iterdir()):
        if (p.name.startswith("_") or p.suffix.lower() not in (".csv", ".xlsx", ".json")
                or p.name in ("mentors.csv", "generation_report.json")):
            continue
        sf = read_any(p)
        out[sf.source] = sf
    return out
